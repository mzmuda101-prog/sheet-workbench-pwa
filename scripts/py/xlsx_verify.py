#!/usr/bin/env python3
"""Niezależna kontrola zapisu .xlsx — drugi parser (openpyxl) zamiast SheetJS.

Po co: testy zapisu (save-roundtrip, save-stress, real-file-regression) sprawdzają wynik
tym samym SheetJS, który bierze udział w zapisie — błąd, który SheetJS „wybacza”, przechodzi.
openpyxl czyta plik po swojemu i porównuje TREŚĆ: wartości, formuły (także rozwinięte
formuły dzielone), typy, style komórek, scalenia, listy (DV), formatowanie warunkowe,
zamrożenia, szerokości kolumn, nazwy, tabele, komentarze, linki.

To NIE jest walidator schematu (Excel może dalej krzyczeć „naprawiono zawartość”, choć
openpyxl plik przeczyta) — sprawdza, czy zapis nie zgubił/nie przekłamał treści.

Polecenia:
  xlsx_verify.py plan ORYGINAŁ              → JSON z cyklami edycji dobranymi do pliku
  xlsx_verify.py verify ORYGINAŁ ZAPIS EDYCJE.json   → porównanie; kod 1 przy błędach
  xlsx_verify.py snapshot PLIK              → zrzut treści (JSON) do podglądu

Format edycji = to, co przyjmuje buildPatchedXlsx (app/xlsx-patch.js), po cyklach:
  [ { "Arkusz": { "B2": {"t": "s", "v": "tekst"}, "C3": null, "D4": {"t": "d", "v": "2026-06-21T00:00:00"} } }, … ]
"""
import datetime as dt
import json
import re
import sys
import warnings
import zipfile
from xml.sax.saxutils import unescape

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.formula import ArrayFormula

# openpyxl ostrzega o rzeczach, których nie porównujemy (np. nagłówek/stopka wydruku) — szum.
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# Te same znaki, które wycina sanitizeXmlText w app/xlsx-patch.js (niedozwolone w XML 1.0).
ILLEGAL_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]|[\ud800-\udfff]")


def load(path):
    # keep_links: zewnętrzne łącza też są treścią; rich_text=False: porównujemy sam tekst.
    return load_workbook(path, data_only=False, keep_links=True, rich_text=False)


# ── zrzut treści ─────────────────────────────────────────────────────────────

def color_key(c):
    if c is None:
        return None
    if c.type == "rgb":
        return c.rgb if isinstance(c.rgb, str) else None
    if c.type == "theme":
        return f"theme{c.theme}{'+%.3f' % c.tint if c.tint else ''}"
    if c.type == "indexed":
        return f"idx{c.indexed}"
    return None


def style_key(cell, with_numfmt=True):
    f, fl, b, a = cell.font, cell.fill, cell.border, cell.alignment
    key = {
        "font": [f.name, f.sz, f.b, f.i, f.u, f.strike, color_key(f.color)],
        "fill": [fl.fill_type, color_key(getattr(fl, "fgColor", None))] if getattr(fl, "fill_type", None) else None,
        "border": [(s.style, color_key(s.color)) if s is not None else None
                   for s in (b.left, b.right, b.top, b.bottom)],
        "align": [a.horizontal, a.vertical, a.wrap_text, a.text_rotation, a.indent],
        "prot": [cell.protection.locked, cell.protection.hidden],
    }
    if with_numfmt:
        key["numfmt"] = cell.number_format
    return key


def value_key(cell):
    v = cell.value
    if isinstance(v, ArrayFormula):
        return {"array": v.text, "ref": v.ref}
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return {"date": v.isoformat()}
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def snapshot_sheet(ws):
    cells = {}
    for (r, c), cell in ws._cells.items():
        cells[f"{get_column_letter(c)}{r}"] = {
            "v": value_key(cell),
            "type": cell.data_type,
            "style": style_key(cell),
            "link": cell.hyperlink.target if cell.hyperlink is not None else None,
            "comment": cell.comment.text if cell.comment is not None else None,
        }
    dvs = sorted(
        json.dumps([dv.type, str(dv.sqref), dv.formula1, dv.formula2, dv.operator, dv.allow_blank,
                    dv.showDropDown, dv.errorStyle], ensure_ascii=False)
        for dv in ws.data_validations.dataValidation
    )
    cf = []
    for rng in ws.conditional_formatting:
        for rule in rng.rules:
            cf.append(json.dumps([str(rng.sqref), rule.type, rule.operator, list(rule.formula or []),
                                  rule.priority, rule.stopIfTrue], ensure_ascii=False))
    return {
        "state": ws.sheet_state,
        "freeze": ws.freeze_panes,
        "merged": sorted(str(m) for m in ws.merged_cells.ranges),
        "dv": dvs,
        "cf": sorted(cf),
        "autofilter": ws.auto_filter.ref,
        "tables": sorted(f"{t.name}@{t.ref}" for t in ws.tables.values()),
        "cols": {k: [d.width, d.hidden, d.outlineLevel] for k, d in sorted(ws.column_dimensions.items())
                 if d.customWidth or d.hidden or d.outlineLevel},
        "rows": {r: [d.ht, d.hidden, d.outlineLevel] for r, d in sorted(ws.row_dimensions.items())
                 if d.ht is not None or d.hidden or d.outlineLevel},
        "print_area": ws.print_area,
        "tab_color": color_key(ws.sheet_properties.tabColor),
        "cells": cells,
    }


def snapshot(path):
    wb = load(path)
    names = {}
    for name, dn in wb.defined_names.items():
        names[name] = dn.attr_text
    for ws in wb.worksheets:
        for name, dn in ws.defined_names.items():
            names[f"{ws.title}!{name}"] = dn.attr_text
    return {
        "sheets": wb.sheetnames,
        "names": names,
        "data": {ws.title: snapshot_sheet(ws) for ws in wb.worksheets},
    }


def default_style(path):
    """Styl komórki bez formatowania w danym pliku — jej usunięcie przy czyszczeniu nic nie gubi."""
    wb = load(path)
    return style_key(wb.worksheets[0].cell(row=1048576, column=16384))


snapshot_default_style = None


# ── porównanie ──────────────────────────────────────────────────────────────

def same_value(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool) and not isinstance(b, bool):
        return abs(a - b) <= 1e-9 * max(1.0, abs(a), abs(b))
    return a == b


def merge_cycles(cycles):
    """Edycje z kolejnych zapisów → stan końcowy (późniejszy cykl wygrywa)."""
    final = {}
    for cycle in cycles:
        for sheet, cells in cycle.items():
            final.setdefault(sheet, {}).update(cells)
    return final


def expected_check(ref, payload, got):
    """Zwraca listę błędów dla komórki edytowanej (got = zrzut komórki albo None)."""
    v = got["v"] if got else None
    if payload is None:
        return [] if v is None else [f"{ref}: miała zostać wyczyszczona, jest {v!r}"]
    t, want = payload["t"], payload["v"]
    if t == "s":
        want = ILLEGAL_XML.sub("", str(want))
        if v != want:
            return [f"{ref}: tekst {v!r} ≠ oczekiwany {want!r}"]
        return []
    if t == "n":
        if not same_value(v, want):
            return [f"{ref}: liczba {v!r} ≠ oczekiwana {want!r}"]
        return []
    if t == "d":
        if not isinstance(v, dict) or "date" not in v:
            return [f"{ref}: data zapisana bez formatu daty (odczyt: {v!r}, format {got and got['style']['numfmt']!r})"]
        if dt.datetime.fromisoformat(v["date"]) != dt.datetime.fromisoformat(want):
            return [f"{ref}: data {v['date']} ≠ oczekiwana {want}"]
        return []
    return [f"{ref}: nieznany typ edycji {t!r}"]


def verify(orig_path, saved_path, cycles):
    global snapshot_default_style
    a, b = snapshot(orig_path), snapshot(saved_path)
    snapshot_default_style = default_style(orig_path)
    edits = merge_cycles(cycles)
    errors, warns = [], []

    if a["sheets"] != b["sheets"]:
        errors.append(f"lista arkuszy: {a['sheets']} → {b['sheets']}")
    if a["names"] != b["names"]:
        errors.append(f"nazwy zdefiniowane zmienione: {a['names']} → {b['names']}")

    for sheet in a["sheets"]:
        sa, sb = a["data"].get(sheet), b["data"].get(sheet)
        if sb is None:
            continue
        pre = f"[{sheet}] "
        for key in ("state", "freeze", "merged", "dv", "cf", "autofilter", "tables", "cols", "print_area", "tab_color"):
            if sa[key] != sb[key]:
                errors.append(f"{pre}{key}: {sa[key]!r} → {sb[key]!r}")
        # Wiersze: porównujemy tylko te z oryginału (nowy wiersz może dostać domyślne atrybuty).
        for r, dims in sa["rows"].items():
            if sb["rows"].get(r) != dims:
                errors.append(f"{pre}wiersz {r}: {dims} → {sb['rows'].get(r)}")

        sheet_edits = edits.get(sheet, {})
        ca, cb = sa["cells"], sb["cells"]
        for ref in sorted(set(ca) | set(cb), key=lambda x: (int(re.sub(r"\D", "", x)), len(x), x)):
            before, after = ca.get(ref), cb.get(ref)
            if ref in sheet_edits:
                payload = sheet_edits[ref]
                errors += [pre + e for e in expected_check(ref, payload, after)]
                if before is not None:
                    if after is None:
                        if before["style"] != snapshot_default_style:
                            errors.append(f"{pre}{ref}: wyczyszczenie usunęło też formatowanie komórki "
                                          f"(Excel przy „Wyczyść zawartość” je zostawia)")
                    else:
                        sb_ = dict(after["style"]); sa_ = dict(before["style"])
                        if payload and payload["t"] == "d":
                            sb_.pop("numfmt"); sa_.pop("numfmt")
                        if sa_ != sb_:
                            errors.append(f"{pre}{ref}: edycja zmieniła styl {sa_} → {sb_}")
                        for k in ("link", "comment"):
                            if before[k] != after[k]:
                                errors.append(f"{pre}{ref}: edycja zmieniła {k}: {before[k]!r} → {after[k]!r}")
                continue
            if before is None:
                errors.append(f"{pre}{ref}: pojawiła się nieedytowana komórka {after['v']!r}")
                continue
            if after is None:
                errors.append(f"{pre}{ref}: zniknęła nieedytowana komórka {before['v']!r}")
                continue
            if not same_value(before["v"], after["v"]):
                errors.append(f"{pre}{ref}: {before['v']!r} → {after['v']!r}")
            elif before["type"] != after["type"]:
                errors.append(f"{pre}{ref}: typ {before['type']} → {after['type']}")
            if before["style"] != after["style"]:
                errors.append(f"{pre}{ref}: styl {before['style']} → {after['style']}")
            for k in ("link", "comment"):
                if before[k] != after[k]:
                    errors.append(f"{pre}{ref}: {k} {before[k]!r} → {after[k]!r}")
    return errors, warns


# ── dobór edycji do pliku ───────────────────────────────────────────────────

SHARED_MASTER = re.compile(r'<c r="([A-Z]+\d+)"[^>]*>\s*<f\b[^>]*\bt="shared"[^>]*\bref="([A-Z0-9:]+)"')


def shared_groups(path):
    """{nazwa arkusza: [(master, zakres)]} — z surowego XML (openpyxl rozwija formuły dzielone)."""
    z = zipfile.ZipFile(path)
    wb = z.read("xl/workbook.xml").decode("utf-8")
    rels = z.read("xl/_rels/workbook.xml.rels").decode("utf-8")
    targets = {m.group(1): m.group(2) for m in re.finditer(r'<Relationship\b[^>]*\bId="([^"]+)"[^>]*\bTarget="([^"]+)"', rels)}
    targets.update({m.group(2): m.group(1) for m in re.finditer(r'<Relationship\b[^>]*\bTarget="([^"]+)"[^>]*\bId="([^"]+)"', rels)})
    out = {}
    for m in re.finditer(r'<sheet\b[^>]*>', wb):
        tag = m.group(0)
        name = unescape(re.search(r'\bname="([^"]*)"', tag).group(1), {"&quot;": '"', "&apos;": "'"})
        rid = re.search(r'\br:id="([^"]+)"', tag).group(1)
        target = targets.get(rid, "")
        part = target.lstrip("/") if target.startswith("/") else "xl/" + target
        if part in z.namelist():
            out[name] = SHARED_MASTER.findall(z.read(part).decode("utf-8"))
    return out


def in_table_header(ws, r, c):
    for t in ws.tables.values():
        c1, r1, c2, _ = range_boundaries(t.ref)
        if r == r1 and c1 <= c <= c2:
            return True
    return False


def plan(path):
    """Dwa cykle zapisu na 1–3 arkuszach z danymi: tekst, liczba, nowe komórki, data,
    czyszczenie, znak sterujący, a gdy są formuły dzielone — edycja mastera i zależnej."""
    wb = load(path)
    groups = shared_groups(path)
    merged_cells = {ws.title: {(r, c) for m in ws.merged_cells.ranges for r, c in m.cells} for ws in wb.worksheets}
    cycle1, cycle2 = {}, {}
    # Pierwsze dwa widoczne arkusze z danymi + pierwszy z formułami dzielonymi (jeśli inny).
    visible = [ws for ws in wb.worksheets if ws.sheet_state == "visible" and ws.max_row >= 2]
    chosen = visible[:2]
    with_shared = next((ws for ws in visible if groups.get(ws.title)), None)
    if with_shared is not None and with_shared not in chosen:
        chosen.append(with_shared)
    for ws in chosen:
        taken = set()

        def pick(pred):
            for (r, c), cell in sorted(ws._cells.items()):
                ref = cell.coordinate
                if r < 2 or ref in taken or (r, c) in merged_cells[ws.title] or in_table_header(ws, r, c):
                    continue
                if pred(cell):
                    taken.add(ref)
                    return ref
            return None

        is_formula = lambda cell: cell.data_type == "f" or isinstance(cell.value, ArrayFormula)
        txt = pick(lambda cell: cell.data_type == "s" and cell.value)
        num = pick(lambda cell: cell.data_type == "n" and isinstance(cell.value, (int, float)) and not is_formula(cell))
        clr = pick(lambda cell: cell.value is not None and not is_formula(cell))
        e1, e2 = {}, {}
        if txt:
            e1[txt] = {"t": "s", "v": 'Zmiana ąęłńóśźż „cudzysłów” & <znak> 😀🇵🇱'}
            e2[txt] = {"t": "s", "v": "druga zmiana tej samej komórki"}
        if num:
            e1[num] = {"t": "n", "v": 12345.678}
        if clr:
            e2[clr] = None
        last_row, last_col = ws.max_row, ws.max_column
        e1[f"A{last_row + 2}"] = {"t": "s", "v": "nowy wiersz pod danymi"}
        e1[f"{get_column_letter(last_col + 1)}2"] = {"t": "n", "v": 7}
        e2[f"B{last_row + 3}"] = {"t": "d", "v": "2026-06-21T00:00:00"}
        e2[f"C{last_row + 3}"] = {"t": "s", "v": "znak\x01sterujący\x0bwycięty"}
        for master, rng in groups.get(ws.title, [])[:1]:
            e1[master] = {"t": "n", "v": 42}
            c1, r1, c2, r2 = range_boundaries(rng)
            dep = next((f"{get_column_letter(c)}{r}" for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)
                        if f"{get_column_letter(c)}{r}" != master), None)
            if dep:
                e2[dep] = {"t": "s", "v": "tekst zamiast formuły"}
        cycle1[ws.title], cycle2[ws.title] = e1, e2
    return [cycle1, cycle2]


def main(argv):
    if len(argv) >= 2 and argv[0] == "plan":
        print(json.dumps(plan(argv[1]), ensure_ascii=False))
        return 0
    if len(argv) >= 2 and argv[0] == "snapshot":
        print(json.dumps(snapshot(argv[1]), ensure_ascii=False, indent=1, default=str))
        return 0
    if len(argv) >= 4 and argv[0] == "verify":
        with open(argv[3], encoding="utf-8") as fh:
            cycles = json.load(fh)
        errors, warns = verify(argv[1], argv[2], cycles)
        for w in warns:
            print("  ⚠️  " + w)
        for e in errors[:40]:
            print("  ❌ " + e)
        if len(errors) > 40:
            print(f"  … i {len(errors) - 40} więcej")
        return 1 if errors else 0
    print(__doc__)
    return 2


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
