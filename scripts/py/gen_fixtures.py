#!/usr/bin/env python3
"""Generator syntetycznych plików testowych .xlsx (openpyxl).

SheetJS w wersji darmowej (scripts/gen-stress-test.js) nie zapisuje stylów, list (DV)
ani formatowania warunkowego — tu to wszystko jest, więc testy zapisu/DV/CF nie muszą
polegać na prywatnych plikach. Dane są zmyślone.

Uwaga: openpyxl zapisuje formuły BEZ wyliczonych wartości (<f> bez <v>) — tak wygląda
plik z wielu generatorów/eksportów, więc to też jest przypadek do obsłużenia.

Użycie: gen_fixtures.py [KATALOG]   (domyślnie scripts/fixtures)
"""
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

OSOBY = ["Anna Nowak", "Łukasz Ąkowski", "Żaneta Źródłowska", "Piotr Śliwiński", "Jan Kowalski"]
STATUSY = ["Nowe", "W toku", "Zamknięte"]


def features(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Dane"
    thin = Side(style="thin", color="FF999999")
    head_fill = PatternFill("solid", fgColor="FFDDEBF7")

    ws["A1"] = "Rejestr zgłoszeń (dane zmyślone)"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:H1")

    heads = ["Nr", "Data", "Osoba", "Status", "Kwota", "Udział", "Opis", "Kwota brutto"]
    for i, h in enumerate(heads, start=1):
        c = ws.cell(row=2, column=i, value=h)
        c.font = Font(bold=True)
        c.fill = head_fill
        c.border = Border(bottom=thin)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    start = dt.datetime(2026, 1, 5)
    for n in range(1, 41):
        r = n + 2
        ws.cell(row=r, column=1, value=n)
        d = ws.cell(row=r, column=2, value=start + dt.timedelta(days=n * 3))
        d.number_format = "yyyy-mm-dd"
        ws.cell(row=r, column=3, value=OSOBY[n % len(OSOBY)])
        ws.cell(row=r, column=4, value=STATUSY[n % len(STATUSY)])
        k = ws.cell(row=r, column=5, value=round(100 + n * 37.25, 2))
        k.number_format = '#,##0.00 "zł"'
        p = ws.cell(row=r, column=6, value=n / 40)
        p.number_format = "0.0%"
        opis = ws.cell(row=r, column=7, value=f"Opis {n}: „cytat” & <znaczek> — długi tekst " + "x" * (n % 7))
        opis.alignment = Alignment(wrap_text=True)
        ws.cell(row=r, column=8, value=f"=E{r}*1.23").number_format = '#,##0.00 "zł"'
        if n % 5 == 0:
            for c in range(1, 9):
                ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="FFFFF2CC")

    ws["J2"] = "Suma"
    ws["J2"].font = Font(bold=True, italic=True, color="FF1F4E79")
    ws["K2"] = "=SUM(E3:E42)"
    ws["J3"] = "Średnia"
    ws["K3"] = "=AVERAGE(Kwoty)"
    ws["J4"] = "Ostatnia data"
    ws["K4"] = "=MAX(B3:B42)"
    ws["K4"].number_format = "dd.mm.yyyy"
    ws["J5"] = "Tablicowa"
    ws["K5"] = "=SUM(E3:E5*F3:F5)"  # zwykła formuła z zakresem (openpyxl nie pisze CSE bez ArrayFormula)
    ws["J6"] = "Link"
    ws["K6"] = "strona projektu"
    ws["K6"].hyperlink = "https://example.com/rejestr"
    ws["J7"] = "Komentarz"
    ws["K7"] = "z notatką"
    ws["K7"].comment = Comment("Notatka testowa — ąęł", "Generator")
    ws["J8"] = "Styl pusty"
    ws["K8"].fill = PatternFill("solid", fgColor="FFE2EFDA")  # pusta komórka tylko ze stylem
    ws["K8"].border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.freeze_panes = "C3"
    ws.auto_filter.ref = "A2:H42"
    for col, w in {"A": 6, "B": 12, "C": 22, "D": 12, "E": 14, "F": 9, "G": 40, "H": 14, "J": 14, "K": 16}.items():
        ws.column_dimensions[col].width = w
    ws.row_dimensions[2].height = 30
    ws.column_dimensions["I"].hidden = True
    ws.sheet_properties.tabColor = "FF70AD47"

    dv_status = DataValidation(type="list", formula1='"Nowe,W toku,Zamknięte"', allow_blank=True)
    dv_status.add("D3:D200")
    dv_osoba = DataValidation(type="list", formula1="=Słowniki!$A$2:$A$6", allow_blank=True,
                              errorStyle="warning", error="Spoza listy", errorTitle="Osoba")
    dv_osoba.add("C3:C200")
    dv_kwota = DataValidation(type="decimal", operator="between", formula1="0", formula2="100000")
    dv_kwota.add("E3:E200")
    for dv in (dv_status, dv_osoba, dv_kwota):
        ws.add_data_validation(dv)

    ws.conditional_formatting.add("E3:E42", CellIsRule(operator="greaterThan", formula=["1000"],
                                                       fill=PatternFill("solid", bgColor="FFF8CBAD")))
    ws.conditional_formatting.add("F3:F42", ColorScaleRule(start_type="min", start_color="FFFFFFFF",
                                                           end_type="max", end_color="FF63BE7B"))
    ws.conditional_formatting.add("A3:H42", FormulaRule(formula=['$D3="Zamknięte"'], font=Font(color="FF808080")))

    sl = wb.create_sheet("Słowniki")
    sl["A1"] = "Osoby"
    sl["B1"] = "Statusy"
    for i, o in enumerate(OSOBY, start=2):
        sl.cell(row=i, column=1, value=o)
    for i, s in enumerate(STATUSY, start=2):
        sl.cell(row=i, column=2, value=s)

    tb = wb.create_sheet("Plan 2026")
    tb.append(["Miesiąc", "Plan", "Wykonanie", "Różnica"])
    for i, m in enumerate(["styczeń", "luty", "marzec", "kwiecień", "maj", "czerwiec"], start=2):
        tb.append([m, 1000 + i * 100, 950 + i * 120, f"=C{i}-B{i}"])
    table = Table(displayName="TabelaPlan", ref="A1:D7")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    tb.add_table(table)

    hid = wb.create_sheet("Ukryty")
    hid["A1"] = "arkusz ukryty"
    hid.sheet_state = "hidden"

    wb.defined_names["Kwoty"] = DefinedName("Kwoty", attr_text="Dane!$E$3:$E$42")
    wb.defined_names["Statusy"] = DefinedName("Statusy", attr_text="'Słowniki'!$B$2:$B$4")
    ws.print_area = "A1:H42"
    wb.save(path)


THREAD_PERSON = "{11111111-1111-1111-1111-111111111111}"


def threaded_comments(path):
    """Komentarz WĄTKOWY jak z Excela 365 (openpyxl ich nie pisze — dokładamy części ręcznie).

    A2 ma wątek (wpis + odpowiedź) i legacy-kopię „[Threaded comment]” w comments, C5 — zwykłą
    notatkę. Test: cell-notes-playwright.js (wątek tylko do odczytu, polskie znaki, zapis go nie rusza).
    """
    import zipfile

    wb = Workbook()
    ws = wb.active
    ws.title = "Dane"
    ws.append(["Nazwa", "Kwota"])
    ws.append(["A", 1])
    ws.append(["B", 2])
    ws["A2"].comment = Comment("[Threaded comment]\n\nComment:\n    Wątek pierwszy", "Jan")
    ws["C5"].comment = Comment("Pusta komórka z notatką", "Ola")
    wb.create_sheet("Drugi")["A1"] = "x"
    tmp = path + ".tmp"
    wb.save(tmp)
    ns = 'xmlns="http://schemas.microsoft.com/office/spreadsheetml/2018/threadedcomments" xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
    threads = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<ThreadedComments {ns}>'
        f'<threadedComment ref="A2" dT="2026-10-19T10:00:00.00" personId="{THREAD_PERSON}" id="{{22222222-2222-2222-2222-222222222222}}"><text>Wątek pierwszy</text></threadedComment>'
        f'<threadedComment ref="A2" dT="2026-10-19T11:00:00.00" personId="{THREAD_PERSON}" id="{{33333333-3333-3333-3333-333333333333}}" parentId="{{22222222-2222-2222-2222-222222222222}}"><text>Odpowiedź</text></threadedComment>'
        "</ThreadedComments>"
    )
    persons = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<personList {ns}>'
        f'<person displayName="Jan" id="{THREAD_PERSON}" userId="jan" providerId="None"/></personList>'
    )
    with zipfile.ZipFile(tmp) as src, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename).decode("utf-8")
            if item.filename == "xl/worksheets/_rels/sheet1.xml.rels":
                data = data.replace("</Relationships>", '<Relationship Id="rId9" Type="http://schemas.microsoft.com/office/2017/10/relationships/threadedComment" Target="../threadedComments/threadedComment1.xml"/></Relationships>')
            elif item.filename == "xl/_rels/workbook.xml.rels":
                data = data.replace("</Relationships>", '<Relationship Id="rId99" Type="http://schemas.microsoft.com/office/2017/10/relationships/person" Target="persons/person.xml"/></Relationships>')
            elif item.filename == "[Content_Types].xml":
                data = data.replace("</Types>", '<Override PartName="/xl/threadedComments/threadedComment1.xml" ContentType="application/vnd.ms-excel.threadedcomments+xml"/><Override PartName="/xl/persons/person.xml" ContentType="application/vnd.ms-excel.person+xml"/></Types>')
            dst.writestr(item, data.encode("utf-8"))
        dst.writestr("xl/threadedComments/threadedComment1.xml", threads.encode("utf-8"))
        dst.writestr("xl/persons/person.xml", persons.encode("utf-8"))
    os.remove(tmp)


def main(argv):
    out_dir = argv[0] if argv else os.path.join(os.path.dirname(__file__), "..", "fixtures")
    os.makedirs(out_dir, exist_ok=True)
    path = os.path.join(out_dir, "openpyxl-features.xlsx")
    features(path)
    print(os.path.normpath(path))
    # bez wypisywania: testy (xlsx-verify) czytają stdout jako JEDNĄ ścieżkę — openpyxl-features
    threaded_comments(os.path.join(out_dir, "threaded-comments.xlsx"))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
